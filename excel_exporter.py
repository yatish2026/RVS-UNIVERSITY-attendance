import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelExporter:
    def __init__(self):
        self.header_font = Font(name='Calibri', size=10, bold=True, color='000000')
        self.title_font = Font(name='Calibri', size=13, bold=True, color='1F497D')
        self.subtitle_font = Font(name='Calibri', size=11, bold=True, color='1F497D')
        self.dept_font = Font(name='Calibri', size=11, bold=True, color='0F243E')
        self.bold_font = Font(name='Calibri', size=9, bold=True)
        self.regular_font = Font(name='Calibri', size=9)
        
        self.header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
        self.dept_fill = PatternFill(start_color='E9EEF4', end_color='E9EEF4', fill_type='solid')
        self.total_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
        self.subtotal_fill = PatternFill(start_color='F2F5F9', end_color='F2F5F9', fill_type='solid')
        self.summary_header_fill = PatternFill(start_color='B4C6E7', end_color='B4C6E7', fill_type='solid')

        self.thin_border = Border(
            left=Side(style='thin', color='D3D3D3'),
            right=Side(style='thin', color='D3D3D3'),
            top=Side(style='thin', color='D3D3D3'),
            bottom=Side(style='thin', color='D3D3D3')
        )
        self.double_bottom_border = Border(
            left=Side(style='thin', color='D3D3D3'),
            right=Side(style='thin', color='D3D3D3'),
            top=Side(style='thin', color='D3D3D3'),
            bottom=Side(style='double', color='000000')
        )

    def export_salary_bill(self, payroll_summary, output_file="Generated_Salary_Bill.xlsx", selected_domain="ALL", selected_department="ALL"):
        wb = openpyxl.Workbook()
        wb.remove(wb.active) # Remove default empty sheet

        all_items = payroll_summary.get("items", [])
        month_year = payroll_summary.get("month_year", "August 2026")
        month_days = payroll_summary.get("month_days", 31)

        # 1. Specific Single Department Export (e.g., IT, CSE, CIVIL, Administration, etc.)
        if selected_department and selected_department != "ALL":
            dept_items = [it for it in all_items if it.get("department") == selected_department]
            if selected_domain and selected_domain != "ALL":
                filtered_by_dom = [it for it in dept_items if it.get("domain") == selected_domain]
                if filtered_by_dom:
                    dept_items = filtered_by_dom
            
            if not dept_items:
                dept_items = all_items

            dept_name = selected_department
            is_teaching = any(it.get("domain") in ["Teaching", "Teaching ID"] for it in dept_items)

            # Sheet 1: Dedicated Department Salary Bill with all staff rows (FIRST & ACTIVE TAB)
            self._create_single_dept_salary_sheet(wb, f"{dept_name[:26]} Salary Bill", dept_name, dept_items, month_year, month_days, is_teaching=is_teaching)

            # Sheet 2: Bank Disbursal for this Department
            self._create_bank_disbursal_sheet(wb, dept_items, f"{dept_name[:18]} - Bank Disbursal", month_year)

            # Sheet 3: Attendance Summary for this Department
            self._create_attendance_summary_sheet(wb, dept_items, f"{dept_name[:18]} - Attendance", month_year, month_days)

            # Sheet 4: Department Overview Summary
            self._create_domain_summary_sheet(wb, dept_name, dept_items, month_year)

            wb.active = 0
            wb.save(output_file)
            print(f"Department Salary Bill successfully exported for {dept_name} to {output_file}")
            return output_file

        # 2. Filter items if specific domain requested
        elif selected_domain and selected_domain != "ALL":
            items = [it for it in all_items if it.get("domain") == selected_domain]
            if not items:
                items = all_items
            is_teaching = (selected_domain == "Teaching" or selected_domain == "Teaching ID")

            # 1. Consolidated Domain Sheet with full employee rows and department subtotals (FIRST TAB)
            self._create_salary_sheet_with_dept_subtotals(
                wb, f"{selected_domain[:20]} Salary Bill", items, month_year, month_days, is_teaching=is_teaching
            )

            # 2. Individual Dedicated Sheets for Each Department in this Domain
            dept_groups = {}
            for it in items:
                d = it.get("department", "General") or "General"
                dept_groups.setdefault(d, []).append(it)

            for dept_name, d_items in sorted(dept_groups.items()):
                clean_title = f"{dept_name[:26]}"
                self._create_single_dept_salary_sheet(wb, clean_title, dept_name, d_items, month_year, month_days, is_teaching=is_teaching)

            # 3. Domain Bank Disbursal Sheet
            self._create_bank_disbursal_sheet(wb, items, f"{selected_domain[:15]} - Bank Disbursal", month_year)

            # 4. Domain Executive Summary Overview
            self._create_domain_summary_sheet(wb, selected_domain, items, month_year)

            # 5. Domain Attendance Summary Sheet
            self._create_attendance_summary_sheet(wb, items, f"{selected_domain[:15]} - Attendance", month_year, month_days)

            wb.active = 0
            wb.save(output_file)
            return output_file

        else:
            # 3. Full Master Workbook
            items = all_items
            teaching_items = [it for it in items if it.get("domain") == "Teaching"]
            if teaching_items:
                self._create_salary_sheet_with_dept_subtotals(wb, "Teaching Faculty", teaching_items, month_year, month_days, is_teaching=True)
                t_depts = {}
                for it in teaching_items:
                    t_depts.setdefault(it.get("department", "General") or "General", []).append(it)
                for dept_name, d_items in sorted(t_depts.items()):
                    sheet_title = f"T - {dept_name[:24]}"
                    self._create_single_dept_salary_sheet(wb, sheet_title, dept_name, d_items, month_year, month_days, is_teaching=True)

            # 2. Teaching (ID / Adjunct) Sheet
            teaching_id_items = [it for it in items if it.get("domain") == "Teaching ID"]
            if teaching_id_items:
                self._create_salary_sheet_with_dept_subtotals(wb, "Teaching (ID)", teaching_id_items, month_year, month_days, is_teaching=True)

            # 3. Non-Teaching Staff Master & Individual Department Sheets
            nt_items = [it for it in items if it.get("domain") == "Non-Teaching"]
            if nt_items:
                self._create_salary_sheet_with_dept_subtotals(wb, "Non-Teaching Staff", nt_items, month_year, month_days, is_teaching=False)
                nt_depts = {}
                for it in nt_items:
                    nt_depts.setdefault(it.get("department", "General") or "General", []).append(it)
                for dept_name, d_items in sorted(nt_depts.items()):
                    sheet_title = f"NT - {dept_name[:23]}"
                    self._create_single_dept_salary_sheet(wb, sheet_title, dept_name, d_items, month_year, month_days, is_teaching=False)

            # 4. Admission Staff Sheet
            adm_items = [it for it in items if it.get("domain") == "Admission"]
            if adm_items:
                self._create_salary_sheet_with_dept_subtotals(wb, "Admission Cell", adm_items, month_year, month_days, is_teaching=False)

            # 5. Support Staff - Security & Attenders
            sec_items = [it for it in items if it.get("domain") == "Support Staff" and any(k in (it.get("department") or "").lower() for k in ["security", "water", "attender"])]
            if sec_items:
                self._create_salary_sheet_with_dept_subtotals(wb, "Security & Attender", sec_items, month_year, month_days, is_teaching=False)

            # 6. Support Staff - Transport & Garden
            trans_items = [it for it in items if it.get("domain") == "Support Staff" and not any(k in (it.get("department") or "").lower() for k in ["security", "water", "attender"])]
            if trans_items:
                self._create_salary_sheet_with_dept_subtotals(wb, "Transport & Garden", trans_items, month_year, month_days, is_teaching=False)

            # 7. Management & Chairman Office
            mgt_items = [it for it in items if it.get("domain") == "Management"]
            if mgt_items:
                self._create_salary_sheet_with_dept_subtotals(wb, "Management Staff", mgt_items, month_year, month_days, is_teaching=False)

            # 8. SBF & Mess Units
            sbf_items = [it for it in items if it.get("domain") in ["SBF Facility", "Hostel & Mess"]]
            if sbf_items:
                self._create_salary_sheet_with_dept_subtotals(wb, "SBF & Sharath Mess", sbf_items, month_year, month_days, is_teaching=False)

            # 9. Master Bank Disbursal
            self._create_bank_disbursal_sheet(wb, items, "Bank Disbursal (PNB)", month_year)

            # 10. Master University Summary
            self._create_summary_sheet(wb, payroll_summary, month_year)

            # 11. Attendance Summary Sheet
            self._create_attendance_summary_sheet(wb, items, "Attendance Summary", month_year, month_days)

        # Set first detailed sheet as active by default so it opens directly onto employee rows
        wb.active = 0
        wb.save(output_file)
        print(f"Salary Bill successfully exported with department-separated sheets to {output_file}")
        return output_file

    def _create_summary_sheet(self, wb, summary, month_year):
        ws = wb.create_sheet(title="Executive Summary")
        
        ws.merge_cells("A1:F1")
        ws["A1"] = "Sri Venkateswara College of Engineering and Technology (SVCET / RVS)"
        ws["A1"].font = self.title_font
        ws["A1"].alignment = Alignment(horizontal="center")

        ws.merge_cells("A2:F2")
        ws["A2"] = f"Executive Payroll & Domain Breakdown Summary - {month_year}"
        ws["A2"].font = self.subtitle_font
        ws["A2"].alignment = Alignment(horizontal="center")

        headers = ["Domain / Wing", "Staff Count", "Total Gross Salary", "Total Deductions", "Net Payable Payout", "% of Total"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx, value=h)
            cell.font = self.header_font
            cell.fill = self.summary_header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.thin_border

        domain_stats = summary.get("domain_stats", {})
        total_gross = summary.get("total_gross", 1.0) or 1.0
        current_row = 5

        for d_name, stat in domain_stats.items():
            pct = (stat["gross"] / total_gross) * 100.0 if total_gross > 0 else 0
            ws.cell(row=current_row, column=1, value=d_name).font = self.bold_font
            ws.cell(row=current_row, column=2, value=stat["count"]).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=3, value=round(stat["gross"], 2))
            ws.cell(row=current_row, column=4, value=round(stat["deductions"], 2))
            ws.cell(row=current_row, column=5, value=round(stat["net"], 2)).font = self.bold_font
            ws.cell(row=current_row, column=6, value=f"{pct:.1f}%").alignment = Alignment(horizontal="center")

            for c in range(1, 7):
                ws.cell(row=current_row, column=c).border = self.thin_border
            current_row += 1

        # Grand Total Row
        ws.cell(row=current_row, column=1, value="TOTAL").font = self.header_font
        ws.cell(row=current_row, column=2, value=summary.get("total_employees", 0)).font = self.header_font
        ws.cell(row=current_row, column=2).alignment = Alignment(horizontal="center")
        ws.cell(row=current_row, column=3, value=round(summary.get("total_gross", 0), 2)).font = self.header_font
        ws.cell(row=current_row, column=4, value=round(summary.get("total_deductions", 0), 2)).font = self.header_font
        ws.cell(row=current_row, column=5, value=round(summary.get("total_net", 0), 2)).font = self.header_font
        ws.cell(row=current_row, column=6, value="100.0%").font = self.header_font
        ws.cell(row=current_row, column=6).alignment = Alignment(horizontal="center")

        for c in range(1, 7):
            cell = ws.cell(row=current_row, column=c)
            cell.fill = self.total_fill
            cell.border = self.double_bottom_border

        # Department Table below
        current_row += 3
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
        ws.cell(row=current_row, column=1, value="Detailed Department-wise Breakdown").font = self.subtitle_font
        current_row += 1

        dept_headers = ["Department", "Domain", "Staff Count", "Gross Salary", "Net Salary"]
        for col_idx, h in enumerate(dept_headers, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=h)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.border = self.thin_border
        current_row += 1

        dept_stats = summary.get("department_stats", {})
        for dept_name, dstat in sorted(dept_stats.items(), key=lambda x: (x[1]["domain"], x[0])):
            ws.cell(row=current_row, column=1, value=dept_name)
            ws.cell(row=current_row, column=2, value=dstat["domain"])
            ws.cell(row=current_row, column=3, value=dstat["count"]).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=4, value=round(dstat["gross"], 2))
            ws.cell(row=current_row, column=5, value=round(dstat["net"], 2))
            for c in range(1, 6):
                ws.cell(row=current_row, column=c).border = self.thin_border
            current_row += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    def _create_domain_summary_sheet(self, wb, domain_name, items, month_year):
        ws = wb.create_sheet(title=f"{domain_name[:20]} Summary")
        
        ws.merge_cells("A1:F1")
        ws["A1"] = f"Sri Venkateswara College of Engineering and Technology ({domain_name.upper()})"
        ws["A1"].font = self.title_font
        ws["A1"].alignment = Alignment(horizontal="center")

        ws.merge_cells("A2:F2")
        ws["A2"] = f"Department-wise Summary Report - {month_year}"
        ws["A2"].font = self.subtitle_font
        ws["A2"].alignment = Alignment(horizontal="center")

        headers = ["Department", "Staff Count", "Total Base CTC", "Gross Salary", "Total Deductions", "Net Salary"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx, value=h)
            cell.font = self.header_font
            cell.fill = self.summary_header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.thin_border

        dept_groups = {}
        for it in items:
            d = it.get("department", "General") or "General"
            dept_groups.setdefault(d, []).append(it)

        current_row = 5
        total_staff = len(items)
        tot_std = sum(it.get("standard_salary", 0) for it in items)
        tot_gross = sum(it.get("gross_total", 0) for it in items)
        tot_ded = sum(it.get("tot_ded", 0) for it in items)
        tot_net = sum(it.get("net_salary", 0) for it in items)

        for d_name, d_list in sorted(dept_groups.items()):
            cnt = len(d_list)
            g_sum = sum(it.get("gross_total", 0) for it in d_list)
            std_sum = sum(it.get("standard_salary", 0) for it in d_list)
            d_sum = sum(it.get("tot_ded", 0) for it in d_list)
            n_sum = sum(it.get("net_salary", 0) for it in d_list)

            ws.cell(row=current_row, column=1, value=d_name).font = self.bold_font
            ws.cell(row=current_row, column=2, value=cnt).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=3, value=round(std_sum, 2))
            ws.cell(row=current_row, column=4, value=round(g_sum, 2))
            ws.cell(row=current_row, column=5, value=round(d_sum, 2))
            ws.cell(row=current_row, column=6, value=round(n_sum, 2)).font = self.bold_font

            for c in range(1, 7):
                ws.cell(row=current_row, column=c).border = self.thin_border
            current_row += 1

        # Grand Total Row
        ws.cell(row=current_row, column=1, value="DOMAIN TOTAL").font = self.header_font
        ws.cell(row=current_row, column=2, value=total_staff).font = self.header_font
        ws.cell(row=current_row, column=2).alignment = Alignment(horizontal="center")
        ws.cell(row=current_row, column=3, value=round(tot_std, 2)).font = self.header_font
        ws.cell(row=current_row, column=4, value=round(tot_gross, 2)).font = self.header_font
        ws.cell(row=current_row, column=5, value=round(tot_ded, 2)).font = self.header_font
        ws.cell(row=current_row, column=6, value=round(tot_net, 2)).font = self.header_font

        for c in range(1, 7):
            cell = ws.cell(row=current_row, column=c)
            cell.fill = self.total_fill
            cell.border = self.double_bottom_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    def _create_salary_sheet_with_dept_subtotals(self, wb, sheet_title, items, month_year, month_days, is_teaching=True):
        # Truncate title if longer than 31 chars (Excel limit)
        clean_title = sheet_title[:31]
        ws = wb.create_sheet(title=clean_title)
        
        ws.merge_cells("A1:AB1")
        ws["A1"] = "Sri Venkateswara College of Engineering and Technology"
        ws["A1"].font = self.title_font
        ws["A1"].alignment = Alignment(horizontal="center")

        ws.merge_cells("A2:AB2")
        ws["A2"] = f"{sheet_title} - Salary Bill for the Month of {month_year}"
        ws["A2"].font = self.subtitle_font
        ws["A2"].alignment = Alignment(horizontal="center")

        headers = [
            "Sl. No.", "Emp ID", "Name of the Staff", "Department", "Designation", "Total Salary", "Consolidated", 
            "Basic", "AGP", "Total", "No of Days", "Basic+AGP", "DA (37.31%)", "HRA (16%)", "Arrears", 
            "FA", "TA/SA", "Gross Total", "EPF", "IT", "PT", "WF", "EB", "Mess", "Bus/others", 
            "Tot Ded", "Net Salary", "PNB Account Number", "IFSC Code"
        ]

        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_idx, value=h)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = self.thin_border

        # Group items by department
        dept_groups = {}
        for it in items:
            dept = it.get("department", "General") or "General"
            dept_groups.setdefault(dept, []).append(it)

        current_row = 4
        dept_subtotal_rows = []

        for dept, dept_items in sorted(dept_groups.items()):
            dept_start_row = current_row + 1
            # Department banner
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers))
            dept_cell = ws.cell(row=current_row, column=1, value=f"DEPARTMENT: {dept.upper()} ({len(dept_items)} Staff Members)")
            dept_cell.font = self.dept_font
            dept_cell.fill = self.dept_fill
            current_row += 1

            d_sl = 1
            for it in dept_items:
                ws.cell(row=current_row, column=1, value=d_sl).alignment = Alignment(horizontal="center")
                ws.cell(row=current_row, column=2, value=it.get("emp_id")).alignment = Alignment(horizontal="center")
                ws.cell(row=current_row, column=3, value=it.get("emp_name"))
                ws.cell(row=current_row, column=4, value=it.get("department"))
                ws.cell(row=current_row, column=5, value=it.get("designation"))
                ws.cell(row=current_row, column=6, value=it.get("standard_salary"))
                ws.cell(row=current_row, column=7, value=it.get("consolidated_salary"))
                ws.cell(row=current_row, column=8, value=it.get("basic"))
                ws.cell(row=current_row, column=9, value=it.get("agp"))
                ws.cell(row=current_row, column=10, value=(it.get("basic", 0) or 0) + (it.get("agp", 0) or 0))
                ws.cell(row=current_row, column=11, value=it.get("total_pay_days")).alignment = Alignment(horizontal="center")
                ws.cell(row=current_row, column=12, value=it.get("basic_agp"))
                ws.cell(row=current_row, column=13, value=it.get("da"))
                ws.cell(row=current_row, column=14, value=it.get("hra"))
                ws.cell(row=current_row, column=15, value=it.get("arrears"))
                ws.cell(row=current_row, column=16, value=it.get("fa"))
                ws.cell(row=current_row, column=17, value=it.get("ta_sa"))
                ws.cell(row=current_row, column=18, value=it.get("gross_total")).font = self.bold_font
                ws.cell(row=current_row, column=19, value=it.get("epf"))
                ws.cell(row=current_row, column=20, value=it.get("it"))
                ws.cell(row=current_row, column=21, value=it.get("pt"))
                ws.cell(row=current_row, column=22, value=it.get("wf"))
                ws.cell(row=current_row, column=23, value=it.get("eb"))
                ws.cell(row=current_row, column=24, value=it.get("mess"))
                ws.cell(row=current_row, column=25, value=it.get("bus"))
                ws.cell(row=current_row, column=26, value=it.get("tot_ded"))
                ws.cell(row=current_row, column=27, value=it.get("net_salary")).font = self.bold_font
                ws.cell(row=current_row, column=28, value=it.get("account_no"))
                ws.cell(row=current_row, column=29, value=it.get("ifsc_code"))

                for c in range(1, len(headers) + 1):
                    cell = ws.cell(row=current_row, column=c)
                    cell.font = self.bold_font if c in [18, 27] else self.regular_font
                    cell.border = self.thin_border

                d_sl += 1
                current_row += 1

            dept_end_row = current_row - 1
            # Department Subtotal Row
            ws.cell(row=current_row, column=3, value=f"Subtotal for {dept}").font = self.bold_font
            for col_idx in [6, 7, 8, 9, 10, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23, 24, 25, 26, 27]:
                col_letter = get_column_letter(col_idx)
                cell = ws.cell(row=current_row, column=col_idx)
                cell.value = f"=SUM({col_letter}{dept_start_row}:{col_letter}{dept_end_row})"
                cell.font = self.bold_font

            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=current_row, column=c)
                cell.fill = self.subtotal_fill
                cell.border = self.thin_border

            dept_subtotal_rows.append(current_row)
            current_row += 2 # gap between depts

        # Grand Total Row for sheet
        ws.cell(row=current_row, column=3, value=f"GRAND TOTAL - {sheet_title}").font = self.header_font
        for col_idx in [6, 7, 8, 9, 10, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23, 24, 25, 26, 27]:
            col_letter = get_column_letter(col_idx)
            cell = ws.cell(row=current_row, column=col_idx)
            # Sum up all department subtotal rows
            formula_parts = [f"{col_letter}{r}" for r in dept_subtotal_rows]
            cell.value = f"=SUM({','.join(formula_parts)})" if formula_parts else "=0"
            cell.font = self.header_font

        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=c)
            cell.fill = self.total_fill
            cell.border = self.double_bottom_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    def _create_single_dept_salary_sheet(self, wb, sheet_title, dept_name, items, month_year, month_days, is_teaching=True):
        clean_title = sheet_title[:31]
        ws = wb.create_sheet(title=clean_title)
        
        ws.merge_cells("A1:AB1")
        ws["A1"] = f"Sri Venkateswara College of Engineering and Technology - Dept of {dept_name.upper()}"
        ws["A1"].font = self.title_font
        ws["A1"].alignment = Alignment(horizontal="center")

        ws.merge_cells("A2:AB2")
        ws["A2"] = f"Monthly Salary Bill for {dept_name} - {month_year} ({len(items)} Staff)"
        ws["A2"].font = self.subtitle_font
        ws["A2"].alignment = Alignment(horizontal="center")

        headers = [
            "Sl. No.", "Emp ID", "Name of the Staff", "Designation", "Total Salary", "Consolidated", 
            "Basic", "AGP", "Total", "No of Days", "Basic+AGP", "DA (37.31%)", "HRA (16%)", "Arrears", 
            "FA", "TA/SA", "Gross Total", "EPF", "IT", "PT", "WF", "EB", "Mess", "Bus/others", 
            "Tot Ded", "Net Salary", "PNB Account Number", "IFSC Code"
        ]

        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_idx, value=h)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = self.thin_border

        current_row = 4
        sl_no = 1
        for it in items:
            ws.cell(row=current_row, column=1, value=sl_no).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=2, value=it.get("emp_id")).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=3, value=it.get("emp_name"))
            ws.cell(row=current_row, column=4, value=it.get("designation"))
            ws.cell(row=current_row, column=5, value=it.get("standard_salary"))
            ws.cell(row=current_row, column=6, value=it.get("consolidated_salary"))
            ws.cell(row=current_row, column=7, value=it.get("basic"))
            ws.cell(row=current_row, column=8, value=it.get("agp"))
            ws.cell(row=current_row, column=9, value=(it.get("basic", 0) or 0) + (it.get("agp", 0) or 0))
            ws.cell(row=current_row, column=10, value=it.get("total_pay_days")).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=11, value=it.get("basic_agp"))
            ws.cell(row=current_row, column=12, value=it.get("da"))
            ws.cell(row=current_row, column=13, value=it.get("hra"))
            ws.cell(row=current_row, column=14, value=it.get("arrears"))
            ws.cell(row=current_row, column=15, value=it.get("fa"))
            ws.cell(row=current_row, column=16, value=it.get("ta_sa"))
            ws.cell(row=current_row, column=17, value=it.get("gross_total")).font = self.bold_font
            ws.cell(row=current_row, column=18, value=it.get("epf"))
            ws.cell(row=current_row, column=19, value=it.get("it"))
            ws.cell(row=current_row, column=20, value=it.get("pt"))
            ws.cell(row=current_row, column=21, value=it.get("wf"))
            ws.cell(row=current_row, column=22, value=it.get("eb"))
            ws.cell(row=current_row, column=23, value=it.get("mess"))
            ws.cell(row=current_row, column=24, value=it.get("bus"))
            ws.cell(row=current_row, column=25, value=it.get("tot_ded"))
            ws.cell(row=current_row, column=26, value=it.get("net_salary")).font = self.bold_font
            ws.cell(row=current_row, column=27, value=it.get("account_no"))
            ws.cell(row=current_row, column=28, value=it.get("ifsc_code"))

            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=current_row, column=c)
                cell.font = self.bold_font if c in [17, 26] else self.regular_font
                cell.border = self.thin_border

            sl_no += 1
            current_row += 1

        # Total Row
        ws.cell(row=current_row, column=3, value=f"Total for {dept_name}").font = self.header_font
        for col_idx in [5, 6, 7, 8, 9, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23, 24, 25, 26]:
            col_letter = get_column_letter(col_idx)
            cell = ws.cell(row=current_row, column=col_idx)
            cell.value = f"=SUM({col_letter}4:{col_letter}{current_row-1})"
            cell.font = self.header_font

        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=c)
            cell.fill = self.total_fill
            cell.border = self.double_bottom_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    def _create_bank_disbursal_sheet(self, wb, items, sheet_title, month_year):
        clean_title = sheet_title[:31]
        ws = wb.create_sheet(title=clean_title)
        
        ws.merge_cells("A1:I1")
        ws["A1"] = f"Sri Venkateswara College Of Engineering and Technology - Bank Disbursal ({month_year})"
        ws["A1"].font = self.title_font
        ws["A1"].alignment = Alignment(horizontal="center")

        headers = ["S.No", "Emp ID", "Name of the Staff", "Domain", "Department", "Designation", "Total Salary", "Net Salary", "PNB Account Number", "IFSC Code"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_idx, value=h)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.thin_border

        current_row = 4
        sl_no = 1
        for it in items:
            ws.cell(row=current_row, column=1, value=sl_no).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=2, value=it.get("emp_id")).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=3, value=it.get("emp_name"))
            ws.cell(row=current_row, column=4, value=it.get("domain"))
            ws.cell(row=current_row, column=5, value=it.get("department"))
            ws.cell(row=current_row, column=6, value=it.get("designation"))
            ws.cell(row=current_row, column=7, value=it.get("standard_salary"))
            ws.cell(row=current_row, column=8, value=it.get("net_salary")).font = self.bold_font
            ws.cell(row=current_row, column=9, value=it.get("account_no"))
            ws.cell(row=current_row, column=10, value=it.get("ifsc_code"))

            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=current_row, column=c)
                cell.font = self.bold_font if c == 8 else self.regular_font
                cell.border = self.thin_border

            sl_no += 1
            current_row += 1

        # Total row
        ws.cell(row=current_row, column=3, value="TOTAL DISBURSAL AMOUNT").font = self.header_font
        ws.cell(row=current_row, column=7, value=f"=SUM(G4:G{current_row-1})").font = self.header_font
        ws.cell(row=current_row, column=8, value=f"=SUM(H4:H{current_row-1})").font = self.header_font

        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=c)
            cell.fill = self.total_fill
            cell.border = self.double_bottom_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    def _create_attendance_summary_sheet(self, wb, items, sheet_title, month_year, month_days):
        clean_title = sheet_title[:31]
        ws = wb.create_sheet(title=clean_title)
        
        ws.merge_cells("A1:J1")
        ws["A1"] = f"Monthly Status Report (Summary Report) - {month_year}"
        ws["A1"].font = self.title_font
        ws["A1"].alignment = Alignment(horizontal="center")

        headers = ["S.No", "Emp ID", "Employee Name", "Domain", "Department", "Designation", "Biometric Days", "Holiday", "Availed Leaves", "Total Pay Days"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_idx, value=h)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.thin_border

        current_row = 4
        sl_no = 1
        for it in items:
            ws.cell(row=current_row, column=1, value=sl_no).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=2, value=it.get("emp_id")).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=3, value=it.get("emp_name"))
            ws.cell(row=current_row, column=4, value=it.get("domain"))
            ws.cell(row=current_row, column=5, value=it.get("department"))
            ws.cell(row=current_row, column=6, value=it.get("designation"))
            ws.cell(row=current_row, column=7, value=it.get("biometric_days")).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=8, value=it.get("holiday_days")).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=9, value=it.get("availed_leaves")).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=10, value=it.get("total_pay_days")).alignment = Alignment(horizontal="center")

            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=current_row, column=c)
                cell.font = self.regular_font
                cell.border = self.thin_border

            sl_no += 1
            current_row += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
