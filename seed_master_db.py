import sys
import os
import openpyxl
import xlrd
import sqlite3
import json
import re

sys.stdout.reconfigure(line_buffering=True)

DB_PATH = "payroll_master.db"

def clean_str(val):
    if val is None:
        return ""
    s = str(val).strip()
    if s.endswith('.0') and s[:-2].isdigit():
        return s[:-2]
    return s

def clean_num(val):
    if val is None or val == "":
        return 0.0
    try:
        if isinstance(val, str):
            val = val.replace(",", "").strip()
        return float(val)
    except:
        return 0.0

def init_database():
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    
    cur.execute("DROP TABLE IF EXISTS employees")
    cur.execute("DROP TABLE IF EXISTS monthly_attendance")
    cur.execute("DROP TABLE IF EXISTS monthly_deductions")
    cur.execute("DROP TABLE IF EXISTS payroll_runs")
    cur.execute("DROP TABLE IF EXISTS payroll_items")
    
    cur.execute("""
    CREATE TABLE IF NOT EXISTS employees (
        emp_id TEXT PRIMARY KEY,
        emp_name TEXT NOT NULL,
        domain TEXT NOT NULL DEFAULT 'Teaching',
        category TEXT NOT NULL DEFAULT 'Teaching',
        department TEXT,
        designation TEXT,
        salary_type TEXT DEFAULT 'REGULAR',
        standard_salary REAL DEFAULT 0,
        consolidated_salary REAL DEFAULT 0,
        basic REAL DEFAULT 0,
        agp REAL DEFAULT 0,
        fa REAL DEFAULT 0,
        ta_sa REAL DEFAULT 0,
        epf_fixed REAL DEFAULT 0,
        it_fixed REAL DEFAULT 0,
        account_no TEXT,
        ifsc_code TEXT,
        status TEXT DEFAULT 'ACTIVE',
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    )
    """)
    
    cur.execute("""
    CREATE TABLE IF NOT EXISTS monthly_attendance (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        month_year TEXT NOT NULL,
        emp_id TEXT NOT NULL,
        emp_name TEXT,
        department TEXT,
        domain TEXT,
        biometric_days REAL DEFAULT 0,
        holiday_days REAL DEFAULT 0,
        availed_leaves REAL DEFAULT 0,
        od_days REAL DEFAULT 0,
        total_pay_days REAL DEFAULT 0,
        remarks TEXT,
        UNIQUE(month_year, emp_id)
    )
    """)
    
    cur.execute("""
    CREATE TABLE IF NOT EXISTS monthly_deductions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        month_year TEXT NOT NULL,
        emp_id TEXT NOT NULL,
        mess_deduction REAL DEFAULT 0,
        eb_deduction REAL DEFAULT 0,
        bus_deduction REAL DEFAULT 0,
        income_tax REAL DEFAULT 0,
        arrears REAL DEFAULT 0,
        other_deductions REAL DEFAULT 0,
        remarks TEXT,
        UNIQUE(month_year, emp_id)
    )
    """)
    
    cur.execute("""
    CREATE TABLE IF NOT EXISTS payroll_runs (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        month_year TEXT NOT NULL UNIQUE,
        total_days_in_month INTEGER NOT NULL,
        total_employees INTEGER DEFAULT 0,
        total_gross REAL DEFAULT 0,
        total_deductions REAL DEFAULT 0,
        total_net REAL DEFAULT 0,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    )
    """)
    
    cur.execute("""
    CREATE TABLE IF NOT EXISTS payroll_items (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        run_id INTEGER NOT NULL,
        emp_id TEXT NOT NULL,
        emp_name TEXT,
        domain TEXT,
        category TEXT,
        department TEXT,
        designation TEXT,
        total_salary REAL DEFAULT 0,
        basic REAL DEFAULT 0,
        agp REAL DEFAULT 0,
        month_days INTEGER DEFAULT 31,
        total_pay_days REAL DEFAULT 0,
        basic_agp REAL DEFAULT 0,
        da REAL DEFAULT 0,
        hra REAL DEFAULT 0,
        arrears REAL DEFAULT 0,
        fa REAL DEFAULT 0,
        ta_sa REAL DEFAULT 0,
        gross_total REAL DEFAULT 0,
        epf REAL DEFAULT 0,
        it REAL DEFAULT 0,
        pt REAL DEFAULT 0,
        wf REAL DEFAULT 0,
        eb REAL DEFAULT 0,
        mess REAL DEFAULT 0,
        bus REAL DEFAULT 0,
        tot_ded REAL DEFAULT 0,
        net_salary REAL DEFAULT 0,
        account_no TEXT,
        ifsc_code TEXT,
        FOREIGN KEY (run_id) REFERENCES payroll_runs (id)
    )
    """)
    
    conn.commit()
    conn.close()
    print("Database tables initialized successfully.")

def normalize_key(s):
    return clean_str(s).lower().replace(" ", "").replace(".", "").replace(",", "").replace("-", "")

def build_full_catalog():
    print("Extracting employee catalog across ALL domains and departments...")
    employees = {}

    # 1. Parse raw biometric report to get initial list of active punch IDs
    try:
        wb_bio = xlrd.open_workbook('Daily Attendance Report 01-09-2026 (2).xls')
        for sname in wb_bio.sheet_names():
            sh = wb_bio.sheet_by_name(sname)
            current_dept = "General"
            for r in range(sh.nrows):
                c1 = clean_str(sh.cell_value(r, 1))
                c3 = clean_str(sh.cell_value(r, 3))
                c7 = clean_str(sh.cell_value(r, 7))
                if 'Department:' in c1:
                    current_dept = c3
                if 'Employee Code:' in c1:
                    emp_code = c3
                    emp_name = c7
                    if emp_code and emp_code not in employees:
                        # Determine domain from department name
                        domain = "Non-Teaching"
                        dept_upper = current_dept.upper()
                        if any(d in dept_upper for d in ["CE", "CIVIL", "EEE", "ME", "MECH", "ECE", "CSE", "CSM", "CSD", "CAI", "MBA", "MCA", "HAS", "H&S"]):
                            domain = "Teaching"
                        elif "SECURITY" in dept_upper or "TRANSPORT" in dept_upper or "ATTENDER" in dept_upper or "GARDEN" in dept_upper or "HOUSEKEEPING" in dept_upper:
                            domain = "Support Staff"
                        elif "SBF" in dept_upper:
                            domain = "SBF Facility"
                        elif "SLH" in dept_upper or "MESS" in dept_upper:
                            domain = "Hostel & Mess"

                        employees[emp_code] = {
                            "emp_id": emp_code,
                            "emp_name": emp_name,
                            "domain": domain,
                            "category": domain,
                            "department": current_dept,
                            "designation": "",
                            "salary_type": "REGULAR",
                            "standard_salary": 0.0,
                            "consolidated_salary": 0.0,
                            "basic": 0.0,
                            "agp": 0.0,
                            "fa": 0.0,
                            "ta_sa": 0.0,
                            "epf_fixed": 0.0,
                            "it_fixed": 0.0,
                            "account_no": "",
                            "ifsc_code": ""
                        }
    except Exception as e:
        print("Error reading Biometric report:", e)

    # 2. Extract from No. of Days Report
    try:
        wb_days = xlrd.open_workbook('SVCET July 2026 - No. of Days.xls')
        sh_days = wb_days.sheet_by_name('New')
        current_dept = "General"
        for r in range(sh_days.nrows):
            c1 = clean_str(sh_days.cell_value(r, 1))
            c2 = clean_str(sh_days.cell_value(r, 2))
            c3 = clean_str(sh_days.cell_value(r, 3))
            if 'Department:' in c1:
                current_dept = c2
            elif c1 and c1 not in ["Emp. Code", "Department:", "S.No"]:
                emp_code = c1
                is_teaching = any(t in c3.lower() for t in ["prof", "lecturer", "principal", "hod", "head", "dean"])
                domain = "Teaching" if is_teaching else "Non-Teaching"
                if emp_code not in employees:
                    employees[emp_code] = {
                        "emp_id": emp_code,
                        "emp_name": c2,
                        "domain": domain,
                        "category": domain,
                        "department": current_dept,
                        "designation": c3,
                        "salary_type": "REGULAR",
                        "standard_salary": 0.0,
                        "consolidated_salary": 0.0,
                        "basic": 0.0,
                        "agp": 0.0,
                        "fa": 0.0,
                        "ta_sa": 0.0,
                        "epf_fixed": 0.0,
                        "it_fixed": 0.0,
                        "account_no": "",
                        "ifsc_code": ""
                    }
                else:
                    if c2: employees[emp_code]["emp_name"] = c2
                    if c3: employees[emp_code]["designation"] = c3
                    if current_dept and current_dept != "General": employees[emp_code]["department"] = current_dept
                    employees[emp_code]["domain"] = domain
                    employees[emp_code]["category"] = domain
    except Exception as e:
        print("Error reading No. of Days:", e)

    # Helper name matcher
    name_to_id = {}
    for eid, emp in employees.items():
        k = normalize_key(emp["emp_name"])
        if k:
            name_to_id[k] = eid

    def find_or_create_emp(name_str, default_domain="Non-Teaching", prefix="EMP"):
        k = normalize_key(name_str)
        if k in name_to_id:
            return name_to_id[k]
        for ek, eid in name_to_id.items():
            if (len(k) > 4 and k in ek) or (len(ek) > 4 and ek in k):
                return eid
        new_id = f"{prefix}_{len(employees)+1}"
        employees[new_id] = {
            "emp_id": new_id,
            "emp_name": name_str,
            "domain": default_domain,
            "category": default_domain,
            "department": "General",
            "designation": "",
            "salary_type": "REGULAR",
            "standard_salary": 0.0,
            "consolidated_salary": 0.0,
            "basic": 0.0,
            "agp": 0.0,
            "fa": 0.0,
            "ta_sa": 0.0,
            "epf_fixed": 0.0,
            "it_fixed": 0.0,
            "account_no": "",
            "ifsc_code": ""
        }
        name_to_id[k] = new_id
        return new_id

    # 3. Teaching Faculty Sheet
    try:
        wb = openpyxl.load_workbook('SVCET July2026 - Teaching Staff (2).xlsx', data_only=True)
        sh = wb['Teaching']
        current_dept = "Teaching"
        for r in range(4, sh.max_row + 1):
            c2 = clean_str(sh.cell(r, 2).value)
            c3 = clean_str(sh.cell(r, 3).value)
            c4 = clean_num(sh.cell(r, 4).value)
            c5 = clean_num(sh.cell(r, 5).value)
            c6 = clean_num(sh.cell(r, 6).value)
            c7 = clean_num(sh.cell(r, 7).value)
            c14 = clean_num(sh.cell(r, 14).value)
            c15 = clean_num(sh.cell(r, 15).value)
            c17 = clean_num(sh.cell(r, 17).value)
            c18 = clean_num(sh.cell(r, 18).value)
            c26 = clean_str(sh.cell(r, 26).value)
            c27 = clean_str(sh.cell(r, 27).value)
            
            if not c2 or "TEACHING STAFF" in c2 or "Total" in c2 or "Grand Total" in c2:
                continue
            if "Dept" in c2 or any(d in c2.upper() for d in ["CIVIL", "EEE", "MECH", "ECE", "CSE", "CSM", "CSD", "CAI", "MBA", "MCA", "H&S"]):
                if not c3:
                    current_dept = c2.replace("Dept. of", "").replace("Dept of", "").strip()
                    continue
            if c3:
                eid = find_or_create_emp(c2, default_domain="Teaching", prefix="T")
                employees[eid].update({
                    "emp_name": c2,
                    "domain": "Teaching",
                    "category": "Teaching",
                    "department": current_dept,
                    "designation": c3,
                    "standard_salary": c4,
                    "consolidated_salary": c5,
                    "basic": c6,
                    "agp": c7,
                    "fa": c14,
                    "ta_sa": c15,
                    "epf_fixed": c17,
                    "it_fixed": c18,
                    "account_no": c26,
                    "ifsc_code": c27
                })
    except Exception as e:
        print("Error parsing Teaching Staff:", e)

    # 4. Teaching Faculty (ID) Sheet
    try:
        wb = openpyxl.load_workbook('SVCET July2026 - Teaching Staff (ID) (1).xlsx', data_only=True)
        sh = wb['Teaching']
        current_dept = "Teaching ID"
        for r in range(4, sh.max_row + 1):
            c2 = clean_str(sh.cell(r, 2).value)
            c3 = clean_str(sh.cell(r, 3).value)
            c4 = clean_num(sh.cell(r, 4).value)
            c5 = clean_num(sh.cell(r, 5).value)
            c6 = clean_num(sh.cell(r, 6).value)
            c7 = clean_num(sh.cell(r, 7).value)
            if not c2 or "TEACHING" in c2 or "Total" in c2:
                continue
            if "Dept" in c2:
                current_dept = c2.replace("Dept. of", "").replace("Dept of", "").strip()
                continue
            if c3:
                eid = find_or_create_emp(c2, default_domain="Teaching ID", prefix="TID")
                employees[eid].update({
                    "emp_name": c2,
                    "domain": "Teaching ID",
                    "category": "Teaching ID",
                    "department": current_dept,
                    "designation": c3,
                    "standard_salary": c4,
                    "consolidated_salary": c5,
                    "basic": c6,
                    "agp": c7
                })
    except Exception as e:
        print("Error parsing Teaching (ID):", e)

    # 5. Non-Teaching Staff Sheet
    try:
        wb = openpyxl.load_workbook('SVCET July 2026 - NT Staff Salary Bill (1).xlsx', data_only=True)
        sh = wb['Non Teaching']
        current_dept = "Administrative Office"
        for r in range(4, sh.max_row + 1):
            c3 = clean_str(sh.cell(r, 3).value)
            c4 = clean_str(sh.cell(r, 4).value)
            c5 = clean_num(sh.cell(r, 5).value)
            c6 = clean_num(sh.cell(r, 6).value)
            c7 = clean_num(sh.cell(r, 7).value)
            c8 = clean_num(sh.cell(r, 8).value)
            c15 = clean_num(sh.cell(r, 15).value)
            c16 = clean_num(sh.cell(r, 16).value)
            c18 = clean_num(sh.cell(r, 18).value)
            c19 = clean_num(sh.cell(r, 19).value)
            c28 = clean_str(sh.cell(r, 28).value)
            c29 = clean_str(sh.cell(r, 29).value)
            
            if not c3 or "NON-TEACHING" in c3 or "Total" in c3:
                continue
            if not c4:
                current_dept = c3
                continue
            eid = find_or_create_emp(c3, default_domain="Non-Teaching", prefix="NT")
            employees[eid].update({
                "emp_name": c3,
                "domain": "Non-Teaching",
                "category": "Non-Teaching",
                "department": current_dept,
                "designation": c4,
                "standard_salary": c5,
                "consolidated_salary": c6,
                "basic": c7,
                "agp": c8,
                "fa": c15,
                "ta_sa": c16,
                "epf_fixed": c18,
                "it_fixed": c19,
                "account_no": c28,
                "ifsc_code": c29
            })
    except Exception as e:
        print("Error parsing Non-Teaching:", e)

    # 6. Admission Staff Sheet
    try:
        wb = openpyxl.load_workbook('SVCET July 2026 - NT Admission.xlsx', data_only=True)
        sh = wb['Admission']
        for r in range(5, sh.max_row + 1):
            c2 = clean_str(sh.cell(r, 2).value)
            c3 = clean_str(sh.cell(r, 3).value)
            c4 = clean_num(sh.cell(r, 4).value)
            c5 = clean_num(sh.cell(r, 5).value)
            c26 = clean_str(sh.cell(r, 26).value)
            c27 = clean_str(sh.cell(r, 27).value)
            if not c2 or "NON-TEACHING" in c2 or "Total" in c2:
                continue
            if c3:
                eid = find_or_create_emp(c2, default_domain="Admission", prefix="ADM")
                employees[eid].update({
                    "emp_name": c2,
                    "domain": "Admission",
                    "category": "Admission",
                    "department": "Admissions Cell",
                    "designation": c3,
                    "standard_salary": c4,
                    "consolidated_salary": c5,
                    "account_no": c26,
                    "ifsc_code": c27
                })
    except Exception as e:
        print("Error parsing Admission:", e)

    # 7. Security, WS, Attender & Transport
    try:
        wb = openpyxl.load_workbook('SVCET July 2026 - NT GS, WS, Attender  Transport.xlsx', data_only=True)
        for sname in ['Security', ' Attender', 'Garden Staff', 'Transport']:
            if sname in wb.sheetnames:
                sh = wb[sname]
                dept_name = sname.strip()
                for r in range(5, sh.max_row + 1):
                    c2 = clean_str(sh.cell(r, 2).value)
                    c3 = clean_str(sh.cell(r, 3).value)
                    c4 = clean_num(sh.cell(r, 4).value)
                    c5 = clean_num(sh.cell(r, 5).value)
                    c26 = clean_str(sh.cell(r, 26).value)
                    c27 = clean_str(sh.cell(r, 27).value)
                    if not c2 or "Total" in c2 or "NON-TEACHING" in c2:
                        continue
                    if c3:
                        eid = find_or_create_emp(c2, default_domain="Support Staff", prefix="SUP")
                        employees[eid].update({
                            "emp_name": c2,
                            "domain": "Support Staff",
                            "category": "Support Staff",
                            "department": dept_name,
                            "designation": c3,
                            "standard_salary": c4,
                            "consolidated_salary": c5,
                            "account_no": c26,
                            "ifsc_code": c27
                        })
    except Exception as e:
        print("Error parsing Support Staff:", e)

    # 8. Management Staff
    try:
        wb = openpyxl.load_workbook('SVCET July2026 - MGT Staff Salary Bill.xlsx', data_only=True)
        sh = wb['Management']
        for r in range(4, sh.max_row + 1):
            c2 = clean_str(sh.cell(r, 2).value)
            c3 = clean_str(sh.cell(r, 3).value)
            c4 = clean_num(sh.cell(r, 4).value)
            c5 = clean_num(sh.cell(r, 5).value)
            c26 = clean_str(sh.cell(r, 26).value)
            c27 = clean_str(sh.cell(r, 27).value)
            if not c2 or "Total" in c2:
                continue
            if c3:
                eid = find_or_create_emp(c2, default_domain="Management", prefix="MGT")
                employees[eid].update({
                    "emp_name": c2,
                    "domain": "Management",
                    "category": "Management",
                    "department": "Management / Chairman Office",
                    "designation": c3,
                    "standard_salary": c4,
                    "consolidated_salary": c5,
                    "account_no": c26,
                    "ifsc_code": c27
                })
    except Exception as e:
        print("Error parsing Management Staff:", e)

    # 9. SBF Facility Services
    try:
        wb = openpyxl.load_workbook('SBF.xlsx', data_only=True)
        for sname in wb.sheetnames:
            sh = wb[sname]
            dept_title = "SBF - SLH" if "SLH" in sname or "Sheet1" in sname else "SBF - SVCET"
            for r in range(4, sh.max_row + 1):
                c2 = clean_str(sh.cell(r, 2).value) # Biometric ID
                c3 = clean_str(sh.cell(r, 3).value) # Name
                c4 = clean_str(sh.cell(r, 4).value) # Desig
                c5 = clean_num(sh.cell(r, 5).value) # Gross
                if not c3 or "Total" in c3:
                    continue
                eid = c2 if c2 and c2.isdigit() else find_or_create_emp(c3, default_domain="SBF Facility", prefix="SBF")
                if eid not in employees:
                    employees[eid] = {"emp_id": eid}
                employees[eid].update({
                    "emp_id": eid,
                    "emp_name": c3,
                    "domain": "SBF Facility",
                    "category": "SBF Facility",
                    "department": dept_title,
                    "designation": c4,
                    "standard_salary": c5,
                    "consolidated_salary": c5
                })
    except Exception as e:
        print("Error parsing SBF:", e)

    # 10. Sharath Mess SLH
    try:
        wb = openpyxl.load_workbook('SLH July 2026 - Sharath Mess.xlsx', data_only=True)
        sh = wb['Sheet2']
        for r in range(4, sh.max_row + 1):
            c2 = clean_str(sh.cell(r, 2).value) # Name
            c3 = clean_num(sh.cell(r, 3).value) # Salary
            if not c2 or "Total" in c2:
                continue
            eid = find_or_create_emp(c2, default_domain="Hostel & Mess", prefix="MESS")
            employees[eid].update({
                "emp_name": c2,
                "domain": "Hostel & Mess",
                "category": "Hostel & Mess",
                "department": "Sharath Mess SLH",
                "designation": "Kitchen & Service Staff",
                "standard_salary": c3,
                "consolidated_salary": c3
            })
    except Exception as e:
        print("Error parsing Sharath Mess:", e)

    # Save to SQLite database
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    for eid, emp in employees.items():
        cur.execute("""
        INSERT OR REPLACE INTO employees 
        (emp_id, emp_name, domain, category, department, designation, salary_type, standard_salary, consolidated_salary, basic, agp, fa, ta_sa, epf_fixed, it_fixed, account_no, ifsc_code)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            emp.get("emp_id", eid),
            emp.get("emp_name", ""),
            emp.get("domain", "Non-Teaching"),
            emp.get("category", "Non-Teaching"),
            emp.get("department", "General"),
            emp.get("designation", ""),
            emp.get("salary_type", "REGULAR"),
            emp.get("standard_salary", 0.0),
            emp.get("consolidated_salary", 0.0),
            emp.get("basic", 0.0),
            emp.get("agp", 0.0),
            emp.get("fa", 0.0),
            emp.get("ta_sa", 0.0),
            emp.get("epf_fixed", 0.0),
            emp.get("it_fixed", 0.0),
            emp.get("account_no", ""),
            emp.get("ifsc_code", "")
        ))
    conn.commit()
    conn.close()
    print(f"Total employees cataloged across all domains: {len(employees)}")

if __name__ == "__main__":
    init_database()
    build_full_catalog()
